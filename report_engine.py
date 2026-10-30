import logging
import os
import re
import smtplib
import time
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path

import pandas as pd
import pymysql
from dotenv import load_dotenv
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from pymysql.err import OperationalError

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")
REPORT_DIR = BASE_DIR / "reports"
LOG_DIR = BASE_DIR / "logs"
REPORT_DIR.mkdir(exist_ok=True)
LOG_DIR.mkdir(exist_ok=True)
logging.basicConfig(filename=LOG_DIR / "automation.log", level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s", encoding="utf-8")
NETWORK_ERROR_CODES = {2002, 2003, 2006, 2013, 2017, 2026, 2055}


def required_env(name):
    value = os.getenv(name, "").strip()
    if not value:
        raise RuntimeError(f"Missing required setting: {name} in .env")
    return value


def get_db_config():
    return {"host": required_env("DB_HOST"), "port": int(os.getenv("DB_PORT", "3306")), "user": required_env("DB_USER"), "password": required_env("DB_PASSWORD"), "database": required_env("DB_NAME"), "charset": "utf8mb4", "connect_timeout": int(os.getenv("DB_CONNECT_TIMEOUT", "8")), "read_timeout": int(os.getenv("DB_READ_TIMEOUT", "600")), "write_timeout": int(os.getenv("DB_WRITE_TIMEOUT", "60")), "autocommit": True}


def check_database_connection():
    connection = pymysql.connect(**get_db_config())
    try:
        with connection.cursor() as cursor:
            cursor.execute("SELECT 1")
            cursor.fetchone()
    finally:
        connection.close()


def wait_for_database():
    attempts = int(os.getenv("DB_RETRY_ATTEMPTS", "5"))
    delay = int(os.getenv("DB_RETRY_DELAY_SECONDS", "60"))
    for attempt in range(1, attempts + 1):
        try:
            logging.info("Database connectivity check %s/%s", attempt, attempts)
            check_database_connection()
            logging.info("Database connectivity check succeeded.")
            return
        except OperationalError as exc:
            code = exc.args[0] if exc.args else None
            logging.warning("Database attempt %s failed (code=%s): %s", attempt, code, exc)
            if code not in NETWORK_ERROR_CODES:
                raise
        except (TimeoutError, OSError) as exc:
            logging.warning("Database connectivity attempt %s failed: %s", attempt, exc)
        if attempt < attempts:
            logging.info("Database unavailable; retrying in %s seconds.", delay)
            time.sleep(delay)
    raise ConnectionError("MySQL unavailable after all connectivity attempts.")


def fetch_procedure_result_sets(procedure, args):
    if not re.fullmatch(r"[A-Za-z0-9_]+", procedure):
        raise ValueError("Procedure name can contain only letters, numbers, and underscores.")
    wait_for_database()
    results = []
    connection = pymysql.connect(**get_db_config())
    try:
        with connection.cursor() as cursor:
            logging.info("Executing stored procedure: %s (argument count: %d)", procedure, len(args))
            cursor.callproc(procedure, args)
            while True:
                if cursor.description:
                    columns = [column[0] for column in cursor.description]
                    rows = cursor.fetchall()
                    if columns:
                        results.append((columns, rows))
                if not cursor.nextset():
                    break
    finally:
        connection.close()
    if not results:
        raise RuntimeError("Stored procedure returned no tabular result sets.")
    logging.info("Fetched %d tabular result set(s) from %s.", len(results), procedure)
    return results


def safe_sheet_name(name, used):
    cleaned = re.sub(r"[\[\]:/*?\\]", "_", str(name)).strip() or "Report"
    cleaned = cleaned[:31]
    base, counter = cleaned, 2
    while cleaned.lower() in used:
        suffix = f"_{counter}"
        cleaned = base[:31 - len(suffix)] + suffix
        counter += 1
    used.add(cleaned.lower())
    return cleaned


def format_workbook(path):
    from openpyxl import load_workbook
    header_fill = PatternFill("solid", fgColor="174A5B")
    header_font = Font(name="Aptos", bold=True, color="FFFFFF", size=11)
    body_font = Font(name="Aptos", size=10)
    side = Side(style="thin", color="D6DEE3")
    border = Border(left=side, right=side, top=side, bottom=side)
    workbook = load_workbook(path)
    for ws in workbook.worksheets:
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A2"
        if ws.max_row and ws.max_column:
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[1].height = 32
            ws.auto_filter.ref = ws.dimensions
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.border = border
                cell.alignment = Alignment(vertical="center")
        for col_idx in range(1, ws.max_column + 1):
            max_len = 0
            for row_idx in range(1, ws.max_row + 1):
                value = ws.cell(row=row_idx, column=col_idx).value
                if value is not None:
                    max_len = max(max_len, len(str(value)))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 3, 12), 40)
    workbook.save(path)


def create_excel(results, report_label="report"):
    safe_label = re.sub(r"[^A-Za-z0-9_-]+", "_", report_label).strip("_") or "report"
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    path = REPORT_DIR / f"{safe_label}_{timestamp}.xlsx"
    used = set()
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for index, (columns, rows) in enumerate(results, start=1):
            sheet = safe_sheet_name(f"Report_{index}", used)
            pd.DataFrame(rows, columns=columns).to_excel(writer, sheet_name=sheet, index=False)
    format_workbook(path)
    if not path.exists() or path.stat().st_size == 0:
        raise RuntimeError("Excel report was not created correctly.")
    logging.info("Excel report created: %s (%s bytes)", path, path.stat().st_size)
    return path


def send_report_email(report_path, recipients):
    sender = required_env("MAIL_FROM")
    password = required_env("MAIL_APP_PASSWORD")
    host = os.getenv("MAIL_SMTP_HOST", "smtp.yandex.com").strip()
    port = int(os.getenv("MAIL_SMTP_PORT", "465"))
    message = EmailMessage()
    message["From"] = sender
    message["To"] = ", ".join(recipients)
    message["Subject"] = f"Automated Report: {report_path.stem}"
    message.set_content("Hello,\n\nPlease find the requested Excel report attached.\n\nThis report was generated by the report automation.\n")
    with open(report_path, "rb") as file:
        message.add_attachment(file.read(), maintype="application", subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=report_path.name)
    logging.info("Connecting to SMTP server %s:%s", host, port)
    with smtplib.SMTP_SSL(host, port, timeout=30) as server:
        server.login(sender, password)
        server.send_message(message, from_addr=sender, to_addrs=recipients)
    logging.info("SMTP server accepted report email for %d recipient(s).", len(recipients))
